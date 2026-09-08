"""Iter142 P0 · Vehicle-wise Cost Report — Excel factory.
5-sheet workbook: Summary · Expenses · Repairs · By Category · By Month.
Pure projection of cost-summary + repair-history payloads.
"""
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

CURRENCY_FMT = '_-₹* #,##0.00_-;-₹* #,##0.00_-;_-₹* "-"??_-;_-@_-'
BOLD = Font(bold=True)
HDR_FILL = PatternFill("solid", fgColor="F4F4F5")


def _hdr(ws, headers, row=1):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = BOLD
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="left", vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_vehicle_cost_xlsx(company: dict, vehicle: dict,
                             cost_summary: dict, repair_history: dict,
                             period: dict) -> bytes:
    wb = openpyxl.Workbook()

    # 1 · Summary
    ws = wb.active
    ws.title = "Summary"
    is_supplier = (vehicle.get("vehicle_type") or "").lower() == "supplier"
    rows = [
        ("Company",       (company or {}).get("name") or "—"),
        ("Vehicle",       vehicle.get("vehicle_number") or "—"),
        ("Ownership",     "Supplier-owned" if is_supplier else "Company-owned"),
        ("Supplier" if is_supplier else "Owner",
         (vehicle.get("supplier_name") if is_supplier else vehicle.get("owner_name")) or "—"),
        ("Supplier Mobile", vehicle.get("supplier_mobile") or "—") if is_supplier
            else ("Owner Phone", vehicle.get("owner_phone") or "—"),
        ("Make / Model",  vehicle.get("make_model") or "—"),
        ("Capacity (T)",  vehicle.get("capacity_tons") or "—"),
        ("Report From",   period.get("from") or "All"),
        ("Report To",     period.get("to")   or "All"),
        ("Category Filter", period.get("category") or "All"),
        ("", ""),
        ("Total Vehicle Cost", cost_summary.get("total_cost", 0)),
        ("Repair Cost",         cost_summary.get("repair_total", 0)),
        ("Operational Expense", cost_summary.get("non_trip_total", 0)),
        ("Trip-linked Cost",    cost_summary.get("trip_linked_total", 0)),
        ("Expense Count",       cost_summary.get("expense_count", 0)),
    ]
    for i, (k, v) in enumerate(rows, start=1):
        a = ws.cell(row=i, column=1, value=k); a.font = BOLD
        b = ws.cell(row=i, column=2, value=v)
        if isinstance(v, (int, float)) and k not in ("Expense Count", "Capacity (T)"):
            b.number_format = CURRENCY_FMT
    _widths(ws, [26, 40])

    # 2 · Expenses
    ws = wb.create_sheet("Expenses")
    _hdr(ws, ["Date", "Category", "Description", "Vendor", "Trip/Repair", "Amount"])
    r = 2
    for row in (cost_summary.get("rows") or []):
        vendor = row.get("party_name") if row.get("party_type") == "vendor" else ""
        ref = "repair" if row.get("repair_event_id") else ("trip" if row.get("trip_id") else (row.get("source_type") or ""))
        ws.cell(row=r, column=1, value=row.get("date") or "")
        ws.cell(row=r, column=2, value=row.get("category") or "")
        ws.cell(row=r, column=3, value=row.get("narration") or row.get("remarks") or "")
        ws.cell(row=r, column=4, value=vendor or "")
        ws.cell(row=r, column=5, value=ref)
        c = ws.cell(row=r, column=6, value=float(row.get("amount") or 0))
        c.number_format = CURRENCY_FMT
        r += 1
    # total row
    tl = ws.cell(row=r, column=5, value="Total"); tl.font = BOLD
    tv = ws.cell(row=r, column=6, value=float(cost_summary.get("total_cost") or 0))
    tv.font = BOLD; tv.number_format = CURRENCY_FMT
    _widths(ws, [12, 14, 45, 22, 12, 16])

    # 3 · Repairs
    ws = wb.create_sheet("Repairs")
    _hdr(ws, ["Date", "Description", "Workshop", "Vendor Payable",
              "Mechanic Payable", "Repair Cost", "Status"])
    r = 2
    for ev in ((repair_history or {}).get("events") or []):
        ws.cell(row=r, column=1, value=ev.get("event_date") or "")
        ws.cell(row=r, column=2, value=ev.get("description") or "")
        ws.cell(row=r, column=3, value=ev.get("workshop_name") or "")
        for col, k in [(4, "vendor_payable"), (5, "mechanic_payable"), (6, "total_repair_cost")]:
            c = ws.cell(row=r, column=col, value=float(ev.get(k) or 0))
            c.number_format = CURRENCY_FMT
        ws.cell(row=r, column=7, value=(ev.get("status") or "").upper())
        r += 1
    tl = ws.cell(row=r, column=5, value="Total Repair Cost"); tl.font = BOLD
    tv = ws.cell(row=r, column=6, value=float((repair_history or {}).get("total_repair_cost") or 0))
    tv.font = BOLD; tv.number_format = CURRENCY_FMT
    _widths(ws, [12, 30, 22, 16, 16, 14, 12])

    # 4 · By Category
    ws = wb.create_sheet("By Category")
    _hdr(ws, ["Category", "Amount"])
    r = 2
    for c in (cost_summary.get("by_category") or []):
        ws.cell(row=r, column=1, value=c.get("category") or "")
        cc = ws.cell(row=r, column=2, value=float(c.get("amount") or 0))
        cc.number_format = CURRENCY_FMT
        r += 1
    _widths(ws, [22, 16])

    # 5 · By Month
    ws = wb.create_sheet("By Month")
    _hdr(ws, ["Month", "Amount"])
    r = 2
    for m in (cost_summary.get("by_month") or []):
        ws.cell(row=r, column=1, value=m.get("month") or "")
        cc = ws.cell(row=r, column=2, value=float(m.get("amount") or 0))
        cc.number_format = CURRENCY_FMT
        r += 1
    _widths(ws, [14, 16])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
