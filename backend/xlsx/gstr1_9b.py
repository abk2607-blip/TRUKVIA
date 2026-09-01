"""Iter132c C3.2 · GSTR-1 §9B XLSX export.

Pure projection of the C3.1 canonical JSON payload. Never recomputes
anything. Six sheets: Summary, CDNR, CDNUR, B2CS_Adjustments,
Commercial_Notes, Cancelled_After_Export.

Deterministic ordering: rows already sorted by (note_date, note_number)
in the C3.1 endpoint's Mongo cursor.
"""
from io import BytesIO
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


HDR_FONT  = Font(name="Calibri", bold=True, color="FFFFFF")
HDR_FILL  = PatternFill("solid", fgColor="0F172A")
BAND      = PatternFill("solid", fgColor="F1F5F9")
CN_FILL   = PatternFill("solid", fgColor="FEF2F2")   # red-50
DN_FILL   = PatternFill("solid", fgColor="EFF6FF")   # blue-50
GOOD_FILL = PatternFill("solid", fgColor="D1FAE5")
BAD_FILL  = PatternFill("solid", fgColor="FEE2E2")
INFO_FILL = PatternFill("solid", fgColor="FEF3C7")
_thin     = Side(border_style="thin", color="CBD5E1")
BORDER    = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

CURRENCY_FMT = '#,##0.00" ₹"'
NUM_FMT      = "#,##0.00"


def _write_header(ws, headers: list, row: int = 4):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = HDR_FONT; c.fill = HDR_FILL; c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _write_row(ws, row_idx: int, values: list, *,
               tint_fill=None, currency_cols=(), num_cols=()):
    for col, v in enumerate(values, 1):
        c = ws.cell(row=row_idx, column=col, value=v)
        c.border = BORDER
        if tint_fill is not None:
            c.fill = tint_fill
        elif row_idx % 2 == 0:
            c.fill = BAND
        if col in currency_cols:
            c.number_format = CURRENCY_FMT
        elif col in num_cols:
            c.number_format = NUM_FMT
        c.alignment = Alignment(vertical="center", wrap_text=True)


def _apply_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ─── Sheet builders ─────────────────────────────────────────────────────

def _build_summary(ws, company: dict, payload: dict):
    ws.cell(row=1, column=1, value=company.get("name", "")).font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=(
        f"GSTR-1 §9B · CN/DN Register  ·  {payload['month']}  ·  "
        f"GSTIN {payload.get('issuer_gstin') or '—'}  ·  "
        f"generated {datetime.now(timezone.utc).strftime('%d-%b-%Y %H:%M UTC')}"
    )).font = Font(italic=True, color="475569")

    _write_header(ws, ["Bucket", "Note Count", "Total ₹", "Credit ₹", "Debit ₹"], row=4)

    t = payload["totals"]
    rows_out = [
        ("CDNR (Registered B2B)",           t["cdnr"]["note_count"],             t["cdnr"]["val"],             t["cdnr"]["cn"],             t["cdnr"]["dn"]),
        ("CDNUR (Unregistered B2CL)",       t["cdnur"]["note_count"],            t["cdnur"]["val"],            t["cdnur"]["cn"],            t["cdnur"]["dn"]),
        ("B2CS Adjustments (net-of Table 7)", t["b2cs_adjustments"]["note_count"], t["b2cs_adjustments"]["val"], t["b2cs_adjustments"]["cn"], t["b2cs_adjustments"]["dn"]),
        ("Commercial Notes (§34 excluded)", t["commercial_notes"]["note_count"], t["commercial_notes"]["val"], t["commercial_notes"]["cn"], t["commercial_notes"]["dn"]),
        ("Cancelled After Export (§9C due)", t["cancelled_after_export"]["note_count"], "", "", ""),
    ]
    for i, r in enumerate(rows_out, start=5):
        _write_row(ws, i, list(r), currency_cols=(3, 4, 5))

    # Reconciliation block
    recon = payload["reconciliation"]
    r0 = 5 + len(rows_out) + 2
    ws.cell(row=r0, column=1, value="Reconciliation").font = Font(bold=True, size=12)
    ws.cell(row=r0 + 1, column=1, value="Reconciled?")
    rc = ws.cell(row=r0 + 1, column=2, value="YES" if recon.get("reconciled") else "NO")
    rc.fill = GOOD_FILL if recon.get("reconciled") else BAD_FILL
    rc.font = Font(bold=True)

    labels = [
        ("Endpoint · GST-true total",     recon["endpoint_gst_true_total"]),
        ("Ground truth · GST-true total", recon["ground_truth_gst_true"]),
        ("Endpoint · GST-false total",    recon["endpoint_gst_false_total"]),
        ("Ground truth · GST-false total",recon["ground_truth_gst_false"]),
        ("Endpoint · issued row count",   recon["endpoint_row_count"]),
        ("Ground truth · issued count",   recon["ground_truth_issued_count"]),
        ("CN total (all)",                     recon["cn_total"]),
        ("DN total (all)",                     recon["dn_total"]),
    ]
    for i, (lbl, val) in enumerate(labels, start=r0 + 2):
        ws.cell(row=i, column=1, value=lbl)
        c = ws.cell(row=i, column=2, value=val)
        if "count" in lbl:
            c.number_format = "0"
        else:
            c.number_format = CURRENCY_FMT

    # Warnings block
    warns = payload.get("warnings", []) or []
    if warns:
        rw = r0 + 2 + len(labels) + 2
        ws.cell(row=rw, column=1, value="Warnings").font = Font(bold=True, size=12)
        for i, w in enumerate(warns, start=rw + 1):
            c = ws.cell(row=i, column=1, value=w)
            c.fill = INFO_FILL

    _apply_widths(ws, [42, 20, 22, 22, 22])


# Column headers are the LOCKED accountant-facing schema for each sheet.
# Tests reference these lists.

CDNR_HEADERS = [
    "Ctin", "Recipient Name", "Note Number", "Note Date", "Type",
    "Note Value ₹", "POS", "Pre-GST", "Reverse Charge",
    "Note Supply Type", "Invoice Number", "Invoice Date",
    "Rate %", "Taxable Value ₹", "CGST ₹", "SGST ₹",
    "IGST ₹", "Cess ₹", "Statutory Reason",
    "QORVENA Reason Code", "Reason Text", "Warnings",
]
CDNUR_HEADERS = [
    "Type (B2CL/EXP)", "Recipient Name", "Note Number", "Note Date", "Type",
    "Note Value ₹", "POS", "Pre-GST",
    "Invoice Number", "Invoice Date",
    "Rate %", "Taxable Value ₹", "CGST ₹", "SGST ₹",
    "IGST ₹", "Cess ₹", "Statutory Reason",
    "QORVENA Reason Code", "Reason Text", "Warnings",
]
B2CS_HEADERS = [
    "Recipient Name", "Note Number", "Note Date", "Type",
    "Note Value ₹", "POS",
    "Invoice Number", "Invoice Date",
    "Rate %", "Taxable Value ₹", "CGST ₹", "SGST ₹",
    "IGST ₹", "Cess ₹", "Statutory Reason",
    "QORVENA Reason Code", "Reason Text", "Advisory",
]
COMMERCIAL_HEADERS = [
    "Note Number", "Note Date", "Type",
    "Recipient Name", "Recipient GSTIN",
    "Invoice Number", "Invoice Date",
    "Subtotal ₹", "Total Amount ₹",
    "QORVENA Reason", "Reason Text", "Info",
]
CANCELLED_HEADERS = [
    "Note Number", "Note Date", "Type",
    "Recipient Name", "Recipient GSTIN",
    "Invoice Number", "Total Amount ₹",
    "Cancelled At", "Cancelled Reason", "Advisory",
]


def _fill_for(kind: str):
    return CN_FILL if kind == "credit" else DN_FILL


def _build_cdnr(ws, payload):
    ws.cell(row=1, column=1, value=(
        f"CDNR · Registered Recipients (Table 9B) · Period {payload['month']}"
    )).font = Font(bold=True, size=12)
    _write_header(ws, CDNR_HEADERS)
    r = 5
    for grp in payload.get("cdnr", []) or []:
        for nt in grp.get("nt", []) or []:
            itm = (nt.get("itms", [{}])[0].get("itm_det") or {})
            _write_row(ws, r, [
                grp.get("ctin", ""),
                nt.get("customer_name", ""),
                nt.get("nt_num", ""),
                nt.get("nt_dt", ""),
                nt.get("ntty", ""),
                float(nt.get("val", 0) or 0),
                nt.get("pos", ""),
                nt.get("p_gst", "N"),
                nt.get("rchrg", "N"),
                "Regular",
                nt.get("inum", ""),
                nt.get("idt", ""),
                float(itm.get("rt", 0) or 0),
                float(itm.get("txval", 0) or 0),
                float(itm.get("camt", 0) or 0),
                float(itm.get("samt", 0) or 0),
                float(itm.get("iamt", 0) or 0),
                float(itm.get("csamt", 0) or 0),
                nt.get("rsn", ""),
                nt.get("reason_code_qorvena", ""),
                nt.get("reason_text", ""),
                "; ".join(nt.get("_warnings", []) or []),
            ], tint_fill=_fill_for(nt.get("kind")),
               currency_cols=(6, 14, 15, 16, 17, 18), num_cols=(13,))
            r += 1
    _apply_widths(ws, [17, 26, 20, 12, 8, 14, 22, 8, 12, 14, 20, 12, 8, 14, 14, 14, 14, 10, 14, 18, 30, 22])


def _build_cdnur(ws, payload):
    ws.cell(row=1, column=1, value=(
        f"CDNUR · Unregistered Recipients (Table 9B) · Period {payload['month']}"
    )).font = Font(bold=True, size=12)
    _write_header(ws, CDNUR_HEADERS)
    r = 5
    for nt in payload.get("cdnur", []) or []:
        itm = (nt.get("itms", [{}])[0].get("itm_det") or {})
        _write_row(ws, r, [
            nt.get("typ", "B2CL"),
            nt.get("customer_name", ""),
            nt.get("nt_num", ""),
            nt.get("nt_dt", ""),
            nt.get("ntty", ""),
            float(nt.get("val", 0) or 0),
            nt.get("pos", ""),
            nt.get("p_gst", "N"),
            nt.get("inum", ""),
            nt.get("idt", ""),
            float(itm.get("rt", 0) or 0),
            float(itm.get("txval", 0) or 0),
            float(itm.get("camt", 0) or 0),
            float(itm.get("samt", 0) or 0),
            float(itm.get("iamt", 0) or 0),
            float(itm.get("csamt", 0) or 0),
            nt.get("rsn", ""),
            nt.get("reason_code_qorvena", ""),
            nt.get("reason_text", ""),
            "; ".join(nt.get("_warnings", []) or []),
        ], tint_fill=_fill_for(nt.get("kind")),
           currency_cols=(6, 12, 13, 14, 15, 16), num_cols=(11,))
        r += 1
    _apply_widths(ws, [12, 26, 20, 12, 8, 14, 22, 8, 20, 12, 8, 14, 14, 14, 14, 10, 14, 18, 30, 22])


def _build_b2cs(ws, payload):
    ws.cell(row=1, column=1, value=(
        f"B2CS Adjustments · Report NET-OF in Table 7 · Period {payload['month']}"
        f"  ·  Not part of §9B statutorily; surfaced here for audit."
    )).font = Font(italic=True, color="B45309", size=10)
    _write_header(ws, B2CS_HEADERS)
    r = 5
    for nt in payload.get("b2cs_adjustments", []) or []:
        itm = (nt.get("itms", [{}])[0].get("itm_det") or {})
        _write_row(ws, r, [
            nt.get("customer_name", ""),
            nt.get("nt_num", ""),
            nt.get("nt_dt", ""),
            nt.get("ntty", ""),
            float(nt.get("val", 0) or 0),
            nt.get("pos", ""),
            nt.get("inum", ""),
            nt.get("idt", ""),
            float(itm.get("rt", 0) or 0),
            float(itm.get("txval", 0) or 0),
            float(itm.get("camt", 0) or 0),
            float(itm.get("samt", 0) or 0),
            float(itm.get("iamt", 0) or 0),
            float(itm.get("csamt", 0) or 0),
            nt.get("rsn", ""),
            nt.get("reason_code_qorvena", ""),
            nt.get("reason_text", ""),
            "; ".join(nt.get("_warnings", []) or []),
        ], tint_fill=_fill_for(nt.get("kind")),
           currency_cols=(5, 10, 11, 12, 13, 14), num_cols=(9,))
        r += 1
    _apply_widths(ws, [26, 20, 12, 8, 14, 22, 20, 12, 8, 14, 14, 14, 14, 10, 14, 18, 30, 30])


def _build_commercial(ws, payload):
    ws.cell(row=1, column=1, value=(
        f"Commercial / Financial Notes · Excluded from GSTR-1 §9B "
        f"per CGST §34 / §15(3)(b) · Period {payload['month']}"
    )).font = Font(italic=True, color="B45309", size=10)
    _write_header(ws, COMMERCIAL_HEADERS)
    r = 5
    for nt in payload.get("commercial_notes", []) or []:
        _write_row(ws, r, [
            nt.get("note_number", ""),
            nt.get("note_date", ""),
            nt.get("ntty", ""),
            nt.get("customer_name", ""),
            nt.get("customer_gstin", ""),
            nt.get("invoice_number", ""),
            nt.get("invoice_date", ""),
            float(nt.get("subtotal", 0) or 0),
            float(nt.get("total_amount", 0) or 0),
            nt.get("reason_code_qorvena", ""),
            nt.get("reason_text", ""),
            nt.get("_info", ""),
        ], tint_fill=_fill_for(nt.get("kind")), currency_cols=(8, 9))
        r += 1
    _apply_widths(ws, [20, 12, 8, 26, 18, 20, 12, 14, 14, 18, 30, 40])


def _build_cancelled(ws, payload):
    ws.cell(row=1, column=1, value=(
        f"Cancelled After Export · GSTR-1 §9C (CDNRA/CDNURA) amendment due "
        f"· Period {payload['month']}"
    )).font = Font(italic=True, color="B91C1C", size=10)
    _write_header(ws, CANCELLED_HEADERS)
    r = 5
    for nt in payload.get("cancelled_after_export", []) or []:
        _write_row(ws, r, [
            nt.get("note_number", ""),
            nt.get("note_date", ""),
            nt.get("ntty", ""),
            nt.get("customer_name", ""),
            nt.get("customer_gstin", ""),
            nt.get("invoice_number", ""),
            float(nt.get("total_amount", 0) or 0),
            nt.get("cancelled_at", ""),
            nt.get("cancelled_reason", ""),
            nt.get("_advisory", ""),
        ], currency_cols=(7,))
        r += 1
    _apply_widths(ws, [20, 12, 8, 26, 18, 20, 14, 22, 30, 40])


def build_gstr1_9b_xlsx(company: dict, payload: dict) -> bytes:
    """Build the full 6-sheet workbook and return raw bytes."""
    wb = openpyxl.Workbook()
    ws_summary = wb.active
    ws_summary.title = "Summary"

    _build_summary(ws_summary, company, payload)
    _build_cdnr(wb.create_sheet("CDNR"), payload)
    _build_cdnur(wb.create_sheet("CDNUR"), payload)
    _build_b2cs(wb.create_sheet("B2CS_Adjustments"), payload)
    _build_commercial(wb.create_sheet("Commercial_Notes"), payload)
    _build_cancelled(wb.create_sheet("Cancelled_After_Export"), payload)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
