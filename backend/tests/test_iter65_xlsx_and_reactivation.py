"""Iter65 · P3 (XLSX Sample) + P4 (Reactivation Guard >90 days) backend contract."""
import os, io, uuid, zipfile, datetime as dt
import httpx, pytest
from dotenv import load_dotenv
load_dotenv("/app/backend/.env")

BASE = os.environ.get("BACKEND_URL_INTERNAL", "http://localhost:8001")
TOKEN = os.environ["DEMO_TOKEN_VALUE"]
HDR = {"Authorization": f"Bearer {TOKEN}"}


def _cid():
    return httpx.get(f"{BASE}/api/companies", headers=HDR, timeout=60).json()[0]["id"]


def _h():
    return {**HDR, "X-Company-Id": _cid()}


# ============ P3 · XLSX sample template ============
def test_xlsx_sample_download():
    r = httpx.get(f"{BASE}/api/vehicles/bulk-import/sample.xlsx", headers=_h(), timeout=60)
    assert r.status_code == 200
    assert r.content[:2] == b"PK", "Not a ZIP/XLSX magic header"
    assert len(r.content) >= 6 * 1024, f"XLSX too small: {len(r.content)} bytes"
    # Inspect sheets
    z = zipfile.ZipFile(io.BytesIO(r.content))
    names = z.namelist()
    assert any("workbook.xml" in n for n in names)
    # Load with openpyxl to check sheets
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)
    assert "Vehicles" in wb.sheetnames
    assert "Instructions" in wb.sheetnames
    ws = wb["Vehicles"]
    # header + 3 sample rows
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) >= 4
    header = rows[0]
    assert "vehicle_number" in header and "vehicle_type" in header


# ============ P4 · Reactivation Guard (>90 days) ============
@pytest.fixture()
def vehicle_deactivated_100d():
    h = _h()
    unique = uuid.uuid4().hex[:6].upper()
    v = httpx.post(f"{BASE}/api/vehicles", headers=h, json={
        "vehicle_number": f"AP99R{unique}", "vehicle_type": "own",
    }, timeout=60).json()
    # Deactivate with effective_date = today - 100 days
    eff = (dt.date.today() - dt.timedelta(days=100)).isoformat()
    r = httpx.patch(f"{BASE}/api/vehicles/{v['id']}/status", headers=h, json={
        "is_active": False, "reason": "long-term storage 100d", "effective_date": eff,
    }, timeout=60)
    assert r.status_code == 200
    return v, eff


def _find_vehicle(vid):
    all_v = httpx.get(f"{BASE}/api/vehicles?active_only=false", headers=_h(), timeout=60).json()
    return next((x for x in all_v if x["id"] == vid), None)


def test_deactivation_persisted_100d(vehicle_deactivated_100d):
    v, eff = vehicle_deactivated_100d
    got = _find_vehicle(v["id"])
    assert got is not None
    assert got["is_active"] is False
    assert got.get("last_status_change_effective_date") == eff
    assert got.get("last_status_change_action") == "deactivated"


def test_reactivation_succeeds_and_audit_has_two(vehicle_deactivated_100d):
    """Backend does NOT gate reactivation — the UI guard is a soft ack."""
    v, _ = vehicle_deactivated_100d
    h = _h()
    r = httpx.patch(f"{BASE}/api/vehicles/{v['id']}/status", headers=h, json={
        "is_active": True, "reason": "back in service after 100d",
        "effective_date": dt.date.today().isoformat(),
    }, timeout=60)
    assert r.status_code == 200
    aud = httpx.get(f"{BASE}/api/vehicles/{v['id']}/status-audit", headers=h, timeout=60).json()
    actions = [a["action"] for a in aud.get("items", [])]
    assert actions.count("deactivated") >= 1 and actions.count("reactivated") >= 1


def test_short_deactivation_10d_no_ui_guard_needed():
    """The <90d case — backend simply stores the effective_date; UI decides."""
    h = _h()
    unique = uuid.uuid4().hex[:6].upper()
    v = httpx.post(f"{BASE}/api/vehicles", headers=h, json={
        "vehicle_number": f"AP99S{unique}", "vehicle_type": "own",
    }, timeout=60).json()
    eff = (dt.date.today() - dt.timedelta(days=10)).isoformat()
    r = httpx.patch(f"{BASE}/api/vehicles/{v['id']}/status", headers=h, json={
        "is_active": False, "reason": "short stop 10d", "effective_date": eff,
    }, timeout=60)
    assert r.status_code == 200
    got = _find_vehicle(v["id"])
    assert got is not None
    assert got["is_active"] is False
    assert got["last_status_change_effective_date"] == eff
