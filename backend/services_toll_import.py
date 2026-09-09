"""Iter148 P0 · FASTag Toll Import — parsers + preview/commit engine.

TRUKVIA principle:
  ENTER ONCE → CALCULATE ONCE → REFLECT EVERYWHERE → REPORT READY.

One real-world Toll transaction produces exactly ONE canonical Expense
(category="Toll", source_type="fastag_import"). No paired db.toll row.
Vehicle Cost / Expense Register / Trip Cost read canonical only.

Supported providers (both share identical schema, differ only in
Vendor column + description pattern):
  • IDFC · Vendor="IDFC" · Description="Issuer Debit Transaction for toll..."
  • LIVQ · Vendor="LIVQ" · Description="FasTag Toll Payment at <PLAZA>"

Sheet: "Account_Summary" · header row at index 7 · data from index 8.
Columns: Transaction Time | Nature (C/D) | Amount | Description |
         Truck Number | Transaction ID | Opening Balance |
         Closing Balance | Vendor.

Eligible rows: Nature=="Debit" AND Truck Number is non-blank.
Credit rows (wallet recharges) are SKIPPED — not toll expenses.
"""
from __future__ import annotations

import hashlib
import io
import re
from datetime import datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from typing import Any, Optional

from db import db

MAX_ROWS_P0 = 2000
POSSIBLE_DUP_AMOUNT_TOL = 0.02      # ±2 %
POSSIBLE_DUP_DATE_WINDOW_DAYS = 1


def _q2(x) -> float:
    try:
        return float(Decimal(str(x)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
    except Exception:
        return 0.0


def _normalise_reg(s: str) -> str:
    """AP 39 UL 6118 → AP39UL6118. Uppercase, strip whitespace."""
    return re.sub(r"\s+", "", (s or "").upper())


def _parse_dt(cell) -> str:
    """`01 Sep 26 11:51 PM` → `2026-09-01`."""
    if cell is None:
        return ""
    if isinstance(cell, datetime):
        return cell.strftime("%Y-%m-%d")
    s = str(cell).strip()
    for fmt in ("%d %b %y %I:%M %p", "%d %b %Y %I:%M %p", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except Exception:
            continue
    return ""


def _iso_ok(s: str) -> bool:
    try:
        datetime.strptime(s or "", "%Y-%m-%d")
        return True
    except Exception:
        return False


def _extract_plaza(desc: str) -> str:
    """`FasTag Toll Payment at PARANUR TOLL` → `PARANUR TOLL`."""
    if not desc:
        return ""
    m = re.search(r"(?:toll payment at|toll at)\s+(.+)", desc, re.I)
    if m:
        return m.group(1).strip()[:120]
    return ""


def make_source_key(vendor: str, cid: str, txn_id: str) -> str:
    return f"toll:{vendor}:{cid}:{txn_id}"


def detect_source(filename: str, blob: bytes) -> Optional[str]:
    """Sniff sheet + header + Vendor column to identify IDFC / LIVQ."""
    if not blob or blob[:4] != b"PK\x03\x04":
        return None
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(blob), data_only=True)
        if "Account_Summary" not in wb.sheetnames:
            return None
        ws = wb["Account_Summary"]
        vendor_col = None
        rows_seen = 0
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 25:
                break
            if i == 7:  # header row
                headers = [str(c or "").strip().lower() for c in row]
                if "transaction id" not in headers or "truck number" not in headers:
                    return None
                try:
                    vendor_col = headers.index("vendor")
                except ValueError:
                    return None
                continue
            if i > 7 and vendor_col is not None:
                v = str(row[vendor_col] or "").strip().upper()
                if v == "IDFC":
                    return "idfc"
                if v == "LIVQ":
                    return "livq"
                rows_seen += 1
                if rows_seen > 5:
                    break
        return None
    except Exception:
        return None


def parse_fastag(blob: bytes) -> tuple[Optional[str], list[dict]]:
    """Return (vendor, rows). vendor=None → unrecognised."""
    from openpyxl import load_workbook
    if not blob or blob[:4] != b"PK\x03\x04":
        return None, []
    try:
        wb = load_workbook(io.BytesIO(blob), data_only=True)
    except Exception:
        return None, []
    if "Account_Summary" not in wb.sheetnames:
        return None, []
    ws = wb["Account_Summary"]

    all_rows = list(ws.iter_rows(values_only=True))
    if len(all_rows) < 9:
        return None, []
    header = [str(c or "").strip() for c in all_rows[7]]
    hmap = {h.lower(): i for i, h in enumerate(header) if h}
    required = {"transaction time", "nature (c/d)", "amount", "description",
                "truck number", "transaction id", "vendor"}
    if not required.issubset(set(hmap)):
        return None, []

    col_time = hmap["transaction time"]
    col_nature = hmap["nature (c/d)"]
    col_amount = hmap["amount"]
    col_desc = hmap["description"]
    col_truck = hmap["truck number"]
    col_txn = hmap["transaction id"]
    col_vendor = hmap["vendor"]

    detected_vendor: Optional[str] = None
    rows: list[dict] = []
    for r_idx in range(8, len(all_rows)):
        row = all_rows[r_idx]
        if not any(row):
            continue
        vendor_raw = str(row[col_vendor] or "").strip().upper()
        if vendor_raw not in ("IDFC", "LIVQ"):
            continue
        if detected_vendor is None:
            detected_vendor = vendor_raw.lower()
        nature = str(row[col_nature] or "").strip().lower()
        if nature != "debit":
            continue
        truck_raw = str(row[col_truck] or "").strip()
        if not truck_raw:
            continue
        try:
            amount = _q2(row[col_amount])
        except Exception:
            amount = 0.0
        date_iso = _parse_dt(row[col_time])
        txn_id = str(row[col_txn] or "").strip()
        # openpyxl may parse numeric txn IDs to float — normalise "745729331.0"
        if txn_id.endswith(".0"):
            txn_id = txn_id[:-2]
        description = str(row[col_desc] or "").strip()[:400]
        plaza = _extract_plaza(description)
        err = None
        if not date_iso:
            err = "Missing or invalid Transaction Time"
        elif amount <= 0:
            err = "Amount must be > 0"
        elif not txn_id:
            err = "Missing Transaction ID"
        rows.append({
            "source": vendor_raw.lower(),
            "source_txn_ref": txn_id,
            "source_vehicle_ref": _normalise_reg(truck_raw),
            "source_vehicle_raw": truck_raw,
            "date": date_iso,
            "amount": amount,
            "description": description,
            "plaza": plaza,
            "row_index": r_idx + 1,
            "error": err,
        })
    return (detected_vendor, rows)


async def build_preview(uid: str, cid: str, vendor: str,
                        rows: list[dict]) -> dict:
    """Bucket every parsed row. Same 5-bucket contract as Iter147."""
    # Auto-resolve vehicles by exact vehicle_number match (normalised).
    veh_cursor = db.vehicles.find(
        {"user_id": uid, "company_id": cid},
        {"_id": 0, "id": 1, "vehicle_number": 1, "is_active": 1},
    )
    vehicles = {}
    async for v in veh_cursor:
        vehicles[_normalise_reg(v.get("vehicle_number", ""))] = v

    processed: list[dict] = []
    seen_keys: dict[str, int] = {}
    for r in rows:
        enriched = dict(r)
        if r.get("error"):
            enriched["bucket"] = "error"
            processed.append(enriched)
            continue

        # Iter148 P0 · exact_duplicate identity is independent of vehicle
        # resolution. Check the canonical source_key BEFORE the vehicle
        # lookup so re-uploads consistently surface prior imports even
        # when the tenant temporarily has no matching own vehicle.
        skey = make_source_key(vendor, cid, r["source_txn_ref"])
        enriched["source_key"] = skey
        if skey in seen_keys:
            enriched["bucket"] = "exact_duplicate"
            enriched["duplicate_reason"] = f"Duplicate within this file (row {seen_keys[skey]})"
            processed.append(enriched)
            continue
        existing = await db.expenses.find_one(
            {"user_id": uid, "company_id": cid, "source_key": skey,
             "is_deleted": {"$ne": True}, "is_reversed": {"$ne": True}},
            {"_id": 0, "id": 1, "date": 1, "amount": 1, "vehicle_number": 1,
             "source": 1, "source_txn_ref": 1},
        )
        if existing:
            enriched["bucket"] = "exact_duplicate"
            enriched["duplicate_reason"] = f"Already imported (Expense {existing['id']})"
            enriched["existing_expense"] = existing
            processed.append(enriched)
            continue
        seen_keys[skey] = r["row_index"]

        veh = vehicles.get(r["source_vehicle_ref"])
        if not veh:
            enriched["bucket"] = "vehicle_mapping_required"
            enriched["resolved_vehicle_id"] = ""
            enriched["resolved_vehicle_number"] = ""
            processed.append(enriched)
            continue
        enriched["resolved_vehicle_id"] = veh["id"]
        enriched["resolved_vehicle_number"] = veh.get("vehicle_number", "")

        pos = await _scan_possible_duplicates(
            uid, cid,
            vehicle_id=veh["id"], date=r["date"],
            amount=r["amount"], plaza=r["plaza"],
        )
        if pos:
            enriched["bucket"] = "possible_duplicate"
            enriched["possible_matches"] = pos
            processed.append(enriched)
            continue

        enriched["bucket"] = "ready"
        processed.append(enriched)

    counts = {b: 0 for b in ("ready", "vehicle_mapping_required",
                             "possible_duplicate", "exact_duplicate", "error")}
    for p in processed:
        counts[p["bucket"]] += 1
    return {
        "vendor": vendor,
        "total_rows": len(processed),
        "counts": counts,
        "rows": processed,
    }


async def _scan_possible_duplicates(uid: str, cid: str, *, vehicle_id: str,
                                    date: str, amount: float,
                                    plaza: str) -> list[dict]:
    if not vehicle_id or not _iso_ok(date):
        return []
    d = datetime.strptime(date, "%Y-%m-%d")
    d_from = (d - timedelta(days=POSSIBLE_DUP_DATE_WINDOW_DAYS)).strftime("%Y-%m-%d")
    d_to = (d + timedelta(days=POSSIBLE_DUP_DATE_WINDOW_DAYS)).strftime("%Y-%m-%d")

    def _similar(row_amount: float, row_narr: str) -> bool:
        if amount > 0 and row_amount > 0 and \
                abs(amount - row_amount) / amount <= POSSIBLE_DUP_AMOUNT_TOL:
            return True
        if plaza and row_narr and plaza.upper()[:8] in row_narr.upper():
            return True
        return False

    matches: list[dict] = []
    async for c in db.expenses.find({
        "user_id": uid, "company_id": cid, "category": "Toll",
        "vehicle_id": vehicle_id,
        "date": {"$gte": d_from, "$lte": d_to},
        "is_deleted": {"$ne": True}, "is_reversed": {"$ne": True},
    }, {"_id": 0, "user_id": 0}).limit(50):
        if _similar(float(c.get("amount") or 0), str(c.get("narration") or "")):
            matches.append({
                "kind": "canonical_expense",
                "source_label": (c.get("source") or "manual").upper() + " Toll",
                "id": c.get("id", ""), "date": c.get("date", ""),
                "amount": float(c.get("amount") or 0),
                "vehicle_number": c.get("vehicle_number", ""),
                "narration": str(c.get("narration") or "")[:120],
            })
            if len(matches) >= 5:
                return matches

    async for t in db.trips.find({
        "user_id": uid, "company_id": cid, "vehicle_id": vehicle_id,
        "date": {"$gte": d_from, "$lte": d_to},
        "$or": [{"has_canonical_expenses": {"$exists": False}},
                {"has_canonical_expenses": False}],
    }, {"_id": 0, "id": 1, "date": 1, "vehicle_number": 1,
         "expenses": 1, "lr_number": 1}).limit(50):
        toll = float((t.get("expenses") or {}).get("toll") or 0)
        if toll <= 0:
            continue
        if _similar(toll, ""):
            matches.append({
                "kind": "trip_legacy_toll",
                "source_label": "Trip Legacy Toll",
                "id": t.get("id", ""), "date": t.get("date", ""),
                "amount": toll,
                "vehicle_number": t.get("vehicle_number", ""),
                "narration": f"LR {t.get('lr_number', '—')}"[:120],
            })
            if len(matches) >= 5:
                return matches
    return matches


async def commit_rows(uid: str, cid: str, vendor: str,
                      rows_in: list[dict]) -> dict:
    """Persist each provided row as ONE canonical Toll Expense."""
    from fastapi import HTTPException
    from models import Expense, now_utc
    from pymongo.errors import DuplicateKeyError
    from routers.expenses import _validate_and_normalise

    if not isinstance(rows_in, list) or not rows_in:
        raise HTTPException(status_code=400, detail="rows must be a non-empty list")
    if len(rows_in) > MAX_ROWS_P0:
        raise HTTPException(status_code=413,
            detail=f"Too many rows ({len(rows_in)} > {MAX_ROWS_P0}); please split the file.")

    batch_id = "tib_" + hashlib.sha1(
        f"{cid}:{now_utc().isoformat()}:{len(rows_in)}".encode()
    ).hexdigest()[:12]

    created = duplicate = failed = 0
    results: list[dict] = []
    for r in rows_in:
        base = {"row_index": r.get("row_index"),
                "source_txn_ref": r.get("source_txn_ref", "")}
        if r.get("error"):
            failed += 1
            results.append({**base, "status": "failed", "error": r["error"]})
            continue
        vid = r.get("resolved_vehicle_id") or ""
        if not vid:
            failed += 1
            results.append({**base, "status": "failed",
                            "error": "Row is missing resolved_vehicle_id"})
            continue
        veh = await db.vehicles.find_one(
            {"id": vid, "user_id": uid, "company_id": cid},
            {"_id": 0, "vehicle_number": 1},
        )
        if not veh:
            failed += 1
            results.append({**base, "status": "failed",
                            "error": "Vehicle not found in tenant"})
            continue

        amount = float(r.get("amount") or 0)
        date_iso = str(r.get("date") or "")
        if amount <= 0 or not _iso_ok(date_iso):
            failed += 1
            results.append({**base, "status": "failed",
                            "error": "Invalid amount or date"})
            continue

        txn_ref = str(r.get("source_txn_ref") or "").strip()
        if not txn_ref:
            failed += 1
            results.append({**base, "status": "failed",
                            "error": "Missing Transaction ID"})
            continue
        skey = r.get("source_key") or make_source_key(vendor, cid, txn_ref)

        description = str(r.get("description") or "").strip()
        plaza = str(r.get("plaza") or "").strip()
        narration = (plaza or description)[:400]
        remarks = f"Txn: {txn_ref}"
        if description and plaza and description != plaza:
            remarks = f"{remarks} · {description[:200]}"
        remarks = remarks[:400]

        try:
            payload = Expense(
                date=date_iso, category="Toll",
                amount=round(amount, 2),
                narration=narration, remarks=remarks,
                vehicle_id=vid, vehicle_number=veh.get("vehicle_number", ""),
                trip_id="", repair_event_id="",
                party_type="cash", party_id="", party_name="",
                vendor_bill_id="", mechanic_work_order_id="",
                supplier_owned_vehicle=False,
                supplier_settlement_mode="n/a",
                settlement_mode="cash_now",
                source_type="fastag_import",
                source_key=skey,
                source=vendor, source_txn_ref=txn_ref,
                source_trip_id="", file_ids=[],
            )
            payload = await _validate_and_normalise(uid, cid, payload)
        except HTTPException as ex:
            failed += 1
            results.append({**base, "status": "failed", "error": str(ex.detail)})
            continue

        existing = await db.expenses.find_one(
            {"user_id": uid, "company_id": cid, "source_key": skey,
             "is_deleted": {"$ne": True}},
            {"_id": 0, "id": 1},
        )
        if existing:
            duplicate += 1
            results.append({**base, "status": "duplicate",
                            "expense_id": existing["id"]})
            continue

        doc = payload.model_dump()
        doc["user_id"] = uid
        doc["company_id"] = cid
        doc["created_by"] = uid
        doc["created_at"] = now_utc().isoformat()
        doc["is_deleted"] = False
        doc["import_batch_id"] = batch_id
        try:
            await db.expenses.insert_one(doc)
        except DuplicateKeyError:
            duplicate += 1
            results.append({**base, "status": "duplicate"})
            continue
        created += 1
        doc.pop("_id", None); doc.pop("user_id", None)
        results.append({**base, "status": "created", "expense_id": doc["id"]})

    return {
        "batch_id": batch_id, "vendor": vendor,
        "created": created, "duplicate": duplicate, "failed": failed,
        "results": results,
    }
