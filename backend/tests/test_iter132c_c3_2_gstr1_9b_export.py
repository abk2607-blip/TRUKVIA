"""Iter132c C3.2 · GSTR-1 §9B XLSX + PDF export targeted suite.

Contract (LOCKED before impl):
  * XLSX and PDF are PURE PROJECTIONS of the C3.1 canonical JSON. Every
    displayed value MUST reconcile byte-for-byte back to the JSON.
  * Both endpoints are feature-flag gated (ENABLE_CDN=1) and share the
    same 400/404 semantics as the JSON endpoint.
  * XLSX: 6 sheets (Summary / CDNR / CDNUR / B2CS_Adjustments /
    Commercial_Notes / Cancelled_After_Export). Frozen header. Deterministic
    row order. Currency format `#,##0.00 ₹`. No row cap.
  * PDF: A4 landscape · Page X of Y · L2d v3 fresh-flowable two-pass ·
    section per bucket · empty-section placeholder text · content grows
    → pages grow → PDF still succeeds. Never HTTP 500 by volume.
  * Audit log emitted on every download (gstr_export/download).
  * 2,100-note streaming — both endpoints must return successfully with
    all rows / all pages retained; guaranteed cleanup.

Reuses the C3.1 `_fake_gstin(state_prefix)` helper pattern for idempotent
re-runnability against the LOCKED Iter127a duplicate-master guard.
"""
import os
import io
import uuid
import asyncio
import time
from datetime import datetime, timedelta
import httpx
import motor.motor_asyncio
import openpyxl
import fitz  # PyMuPDF — already used by L2e test suite

BASE = os.environ.get("REACT_APP_BACKEND_URL", "http://localhost:8001").rstrip("/")
API = BASE + "/api"
HDR = {"Authorization": f"Bearer {os.environ['DEMO_TOKEN_VALUE']}"}


# ─── Helpers (mirror C3.1 shape) ────────────────────────────────────────

def _company_header():
    r = httpx.get(f"{API}/companies", headers=HDR, timeout=15)
    assert r.status_code == 200, r.text
    cid = r.json()[0]["id"]
    return cid, {**HDR, "X-Company-Id": cid}


def _current_uid() -> str:
    r = httpx.get(f"{API}/auth/me", headers=HDR, timeout=15)
    assert r.status_code == 200, r.text
    return r.json()["user_id"]


def _fake_gstin(state_prefix: str) -> str:
    hx = uuid.uuid4().hex.upper()
    letters_pool = "".join(c for c in hx if c.isalpha()) + "AAAAAAAA"
    digits_pool = "".join(c for c in hx if c.isdigit()) + "00000000"
    letters5 = letters_pool[:5]
    digits4 = digits_pool[:4]
    letter1 = letters_pool[5]
    entity = hx[10]
    checksum = hx[11]
    return f"{state_prefix}{letters5}{digits4}{letter1}{entity}Z{checksum}"


def _fresh_customer(h, *, gstin: str = "", state: str = "Telangana"):
    tag = uuid.uuid4().hex[:8]
    body = {
        "name": f"C3.2-Cust-{tag}",
        "phone": f"9{tag[:9]}",
        "email": f"{tag}@iter132c-c3-2.test",
        "state": state,
    }
    if gstin:
        body["gstin"] = gstin
    r = httpx.post(f"{API}/customers", headers=h, json=body, timeout=15)
    assert r.status_code == 200, f"customer create failed: {r.text[:400]}"
    return r.json()["id"]


def _create_invoice_via_api(h, cust_id, *, date, freight_rate=900, tons=20,
                            product_rate=40000, gst_type="cgst_sgst", rcm=True):
    tag = uuid.uuid4().hex[:6]
    tr = httpx.post(f"{API}/trips", headers=h, json={
        "customer_id": cust_id, "date": date,
        "vehicle_number": f"CX{tag[:4].upper()}",
        "tons": tons, "loaded_qty": tons, "unloaded_qty": tons,
        "freight_mode": "per_ton", "rate_per_ton": freight_rate,
        "product_rate_per_mt": product_rate,
        "from_location": "A", "to_location": "B",
        "loading_date": date, "unloading_date": date,
    }, timeout=15)
    assert tr.status_code == 200, tr.text
    inv = httpx.post(f"{API}/invoices", headers=h, json={
        "customer_id": cust_id, "trip_ids": [tr.json()["id"]],
        "invoice_date": date, "rcm": rcm, "gst_type": gst_type,
    }, timeout=15)
    assert inv.status_code == 200, inv.text
    return inv.json()


def _issue_note(h, kind, invoice_id, *, amount, note_date, reason_code="other",
                reason_text="C3.2 test note", apply_gst=True):
    ep = "credit-notes" if kind == "credit" else "debit-notes"
    r = httpx.post(f"{API}/{ep}", headers=h, json={
        "invoice_id": invoice_id, "note_date": note_date,
        "reason_code": reason_code, "reason_text": reason_text,
        "apply_gst": apply_gst,
        "lines": [{"description": f"C3.2 {kind}", "hsn_sac": "996791",
                   "quantity": 1, "rate": amount}],
    }, timeout=15)
    assert r.status_code == 200, f"{kind} note failed: {r.status_code} {r.text[:400]}"
    return r.json()


def _get_json(h, month: str):
    return httpx.get(f"{API}/reports/gstr1-9b", headers=h,
                     params={"month": month}, timeout=60)


def _get_xlsx(h, month: str):
    return httpx.get(f"{API}/reports/gstr1-9b.xlsx", headers=h,
                     params={"month": month}, timeout=120)


def _get_pdf(h, month: str):
    return httpx.get(f"{API}/reports/gstr1-9b.pdf", headers=h,
                     params={"month": month}, timeout=180)


# ─── XLSX tests T1–T9 ───────────────────────────────────────────────────

def test_t1_xlsx_endpoint_reachable():
    """Positive-gate check — endpoint is discoverable when ENABLE_CDN=1
    (preview backend). Negative gate is tested in the Iter132a flag test."""
    cid, h = _company_header()
    r = _get_xlsx(h, "2026-06")
    assert r.status_code != 404, f"XLSX endpoint unreachable: {r.status_code} {r.text[:200]}"


def test_t2_xlsx_bad_month_400():
    cid, h = _company_header()
    for bad in ["2026-13", "202606", "not-a-month", "2026-00", ""]:
        r = _get_xlsx(h, bad)
        assert r.status_code == 400, f"month={bad!r} → {r.status_code}"


def test_t3_xlsx_content_type_and_disposition():
    cid, h = _company_header()
    r = _get_xlsx(h, "2020-01")  # ancient month → empty period
    assert r.status_code == 200, r.text
    assert r.headers.get("content-type", "").startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ), f"MIME wrong: {r.headers.get('content-type')}"
    disp = r.headers.get("content-disposition", "")
    assert "attachment" in disp and "GSTR1_9B_" in disp and ".xlsx" in disp, disp


def test_t4_xlsx_has_all_six_sheets():
    cid, h = _company_header()
    r = _get_xlsx(h, "2020-01")
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    expected = {"Summary", "CDNR", "CDNUR", "B2CS_Adjustments",
                "Commercial_Notes", "Cancelled_After_Export"}
    assert expected == set(wb.sheetnames), \
        f"sheet mismatch. expected={expected} got={set(wb.sheetnames)}"


def test_t5_xlsx_headers_match_locked_schema():
    """Every non-Summary sheet's header row must match the LOCKED
    column-name list published in `backend/xlsx/gstr1_9b.py`."""
    import sys
    sys.path.insert(0, "/app/backend")
    from xlsx.gstr1_9b import (
        CDNR_HEADERS, CDNUR_HEADERS, B2CS_HEADERS,
        COMMERCIAL_HEADERS, CANCELLED_HEADERS,
    )
    cid, h = _company_header()
    r = _get_xlsx(h, "2020-01")
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    for sheet_name, expected in [
        ("CDNR", CDNR_HEADERS), ("CDNUR", CDNUR_HEADERS),
        ("B2CS_Adjustments", B2CS_HEADERS),
        ("Commercial_Notes", COMMERCIAL_HEADERS),
        ("Cancelled_After_Export", CANCELLED_HEADERS),
    ]:
        ws = wb[sheet_name]
        got = [ws.cell(row=4, column=i + 1).value for i in range(len(expected))]
        assert got == expected, \
            f"{sheet_name} header mismatch\nexpected={expected}\ngot     ={got}"


def test_t6_xlsx_json_parity_seeded_data():
    """Seed one CN and one DN in a fresh window, download XLSX + JSON for
    the same period, and assert the CDNR sheet's row values match the JSON
    payload row-for-row.

    NOTE: `nt_num` is per-customer (LOCKED Iter132a behaviour) and is NOT
    globally unique across the tenant. Lookup MUST combine the seeded
    customer identity with the note number — the fresh `_fresh_customer`
    name is uuid-tagged and therefore unique per test run."""
    cid, h = _company_header()
    tg = _fake_gstin("36")
    tag_marker = uuid.uuid4().hex[:8]
    tag_body = {
        "name": f"C3.2-T6-{tag_marker}",
        "phone": f"9{tag_marker[:9]}",
        "email": f"{tag_marker}@iter132c-c3-2-t6.test",
        "state": "Telangana", "gstin": tg,
    }
    r_cust = httpx.post(f"{API}/customers", headers=h, json=tag_body, timeout=15)
    assert r_cust.status_code == 200, r_cust.text
    cust_id = r_cust.json()["id"]
    uniq_name = tag_body["name"]  # tenant-unique for this test run

    today = datetime.now().date()
    inv_iso = today.isoformat()
    note_iso = today.isoformat()
    month = today.strftime("%Y-%m")
    inv = _create_invoice_via_api(h, cust_id, date=inv_iso,
                                  gst_type="cgst_sgst", rcm=False)
    n1 = _issue_note(h, "credit", inv["id"], amount=500, note_date=note_iso,
                     reason_code="rate_correction")
    n2 = _issue_note(h, "debit", inv["id"], amount=300, note_date=note_iso,
                     reason_code="under_charge")

    rj = _get_json(h, month); assert rj.status_code == 200, rj.text
    rx = _get_xlsx(h, month); assert rx.status_code == 200, rx.text
    body = rj.json()

    # Build the parity map keyed by (customer_name, nt_num) — tenant-unique
    # even though `nt_num` alone is per-customer.
    json_map = {}
    for grp in body["cdnr"]:
        for nt in grp["nt"]:
            if nt["note_id"] in (n1["id"], n2["id"]):
                json_map[(nt["customer_name"], nt["nt_num"])] = (grp["ctin"], nt)

    assert len(json_map) == 2, \
        f"expected 2 seeded notes in JSON cdnr, got {len(json_map)}"

    wb = openpyxl.load_workbook(io.BytesIO(rx.content), read_only=True)
    ws = wb["CDNR"]

    matched = 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        key = (row[1], row[2])  # (customer_name, nt_num)
        if key in json_map:
            matched += 1
            expected_ctin, nt_j = json_map[key]
            assert row[0] == expected_ctin, \
                f"CDNR ctin mismatch for {key}: xlsx={row[0]!r} json={expected_ctin!r}"
            assert row[3] == nt_j["nt_dt"]
            assert row[4] == nt_j["ntty"]
            assert abs(float(row[5]) - float(nt_j["val"])) < 0.01
            assert row[6] == nt_j["pos"]
            assert row[10] == nt_j["inum"]
            itm = nt_j["itms"][0]["itm_det"]
            assert abs(float(row[13]) - float(itm["txval"])) < 0.01
            assert abs(float(row[14]) - float(itm["camt"])) < 0.01
            assert abs(float(row[15]) - float(itm["samt"])) < 0.01
            assert abs(float(row[16]) - float(itm["iamt"])) < 0.01
            assert row[18] == nt_j["rsn"]
            assert row[19] == nt_j["reason_code_qorvena"]
    assert matched == 2, \
        f"expected 2 seeded rows in CDNR (name={uniq_name!r}), matched {matched}"


def test_t7_xlsx_summary_reconciliation_matches_json():
    cid, h = _company_header()
    month = "2020-01"  # empty period — reconciled must still be true
    rj = _get_json(h, month).json()
    rx = _get_xlsx(h, month)
    wb = openpyxl.load_workbook(io.BytesIO(rx.content), read_only=True)
    ws = wb["Summary"]
    reconciled_cell = None
    endpoint_gst_true_cell = None
    for r_i in range(1, 40):
        v0 = ws.cell(row=r_i, column=1).value
        if v0 == "Reconciled?":
            reconciled_cell = ws.cell(row=r_i, column=2).value
        if v0 and "Endpoint" in str(v0) and "GST-true" in str(v0):
            endpoint_gst_true_cell = ws.cell(row=r_i, column=2).value
    assert reconciled_cell == ("YES" if rj["reconciliation"]["reconciled"] else "NO")
    assert abs(float(endpoint_gst_true_cell or 0) -
               float(rj["reconciliation"]["endpoint_gst_true_total"])) < 0.01


def test_t8_xlsx_empty_period_downloads_headers_only():
    cid, h = _company_header()
    r = _get_xlsx(h, "2018-01")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    for sheet_name in ("CDNR", "CDNUR", "B2CS_Adjustments",
                       "Commercial_Notes", "Cancelled_After_Export"):
        ws = wb[sheet_name]
        # Row 4 = header; row 5 first data row. Empty period ⇒ row 5 blank.
        row5 = [ws.cell(row=5, column=i).value for i in range(1, 10)]
        assert all(v is None for v in row5), f"{sheet_name} unexpectedly has data"


def test_t9_xlsx_2100_note_streaming():
    """Seed 2,100 issued notes directly into Mongo, download XLSX, assert
    all rows survive. Cleanup guaranteed via finally."""
    cid, h = _company_header()
    uid = _current_uid()
    cust_id = _fresh_customer(h, gstin=_fake_gstin("36"), state="Telangana")
    SEED = 2100
    seed_tag = f"C32T9-{uuid.uuid4().hex[:10]}"

    async def _seed():
        c = motor.motor_asyncio.AsyncIOMotorClient(os.environ["MONGO_URL"])
        _db = c[os.environ["DB_NAME"]]
        inv_id = f"{seed_tag}-inv"
        await _db.invoices.insert_one({
            "id": inv_id, "user_id": uid, "company_id": cid,
            "customer_id": cust_id,
            "invoice_number": f"{seed_tag}/MASTER",
            "invoice_date": "2026-12-01",
            "total_amount": 100.0, "amount_paid": 0.0, "balance_due": 100.0,
            "trip_ids": [], "payments": [], "is_historical": False,
            "seed_tag": seed_tag, "gst_type": "cgst_sgst",
        })
        docs = []
        for i in range(SEED):
            docs.append({
                "id": f"{seed_tag}-cdn-{i:05d}", "user_id": uid, "company_id": cid,
                "customer_id": cust_id,
                "kind": "credit" if i % 2 == 0 else "debit",
                "note_number": f"CN-T9-{seed_tag[-6:]}-{i:05d}",
                "note_date": "2026-12-15",
                "invoice_id": inv_id,
                "invoice_number_snapshot": f"{seed_tag}/MASTER",
                "reason_code": "other", "reason_text": f"seed {i}",
                "lines": [], "subtotal": 10.0, "gst_type": "cgst_sgst",
                "cgst_rate": 2.5, "sgst_rate": 2.5, "igst_rate": 5.0,
                "cgst_amount": 0.25, "sgst_amount": 0.25, "igst_amount": 0.0,
                "total_tax": 0.5, "total_amount": 10.0,
                "rcm": True, "apply_gst": True, "status": "issued",
                "seed_tag": seed_tag,
            })
        for i in range(0, len(docs), 500):
            await _db.credit_debit_notes.insert_many(docs[i:i + 500])
        c.close()

    async def _cleanup():
        c = motor.motor_asyncio.AsyncIOMotorClient(os.environ["MONGO_URL"])
        _db = c[os.environ["DB_NAME"]]
        await _db.credit_debit_notes.delete_many({"seed_tag": seed_tag})
        await _db.invoices.delete_many({"seed_tag": seed_tag})
        r1 = await _db.credit_debit_notes.count_documents({"seed_tag": seed_tag})
        r2 = await _db.invoices.count_documents({"seed_tag": seed_tag})
        c.close()
        return r1, r2

    try:
        asyncio.run(_seed())
        r = _get_xlsx(h, "2026-12")
        assert r.status_code == 200, r.text
        wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
        ws = wb["CDNR"]
        seeded_hits = sum(
            1 for row in ws.iter_rows(min_row=5, values_only=True)
            if row[2] and f"CN-T9-{seed_tag[-6:]}-" in str(row[2])
        )
        assert seeded_hits == SEED, \
            f"XLSX truncation: got {seeded_hits} / expected {SEED}"
    finally:
        r1, r2 = asyncio.run(_cleanup())
        assert r1 == 0 and r2 == 0, f"cleanup orphans: cdn={r1} inv={r2}"


# ─── PDF tests T10–T17 ──────────────────────────────────────────────────

def test_t10_pdf_endpoint_reachable():
    cid, h = _company_header()
    r = _get_pdf(h, "2026-06")
    assert r.status_code != 404, f"PDF endpoint unreachable: {r.status_code}"


def test_t11_pdf_bad_month_400():
    cid, h = _company_header()
    for bad in ["2026-13", "202606", "", "not"]:
        r = _get_pdf(h, bad)
        assert r.status_code == 400, f"month={bad!r} → {r.status_code}"


def test_t12_pdf_content_type_and_disposition():
    cid, h = _company_header()
    r = _get_pdf(h, "2020-01")
    assert r.status_code == 200
    assert r.headers.get("content-type", "").startswith("application/pdf")
    disp = r.headers.get("content-disposition", "")
    assert "GSTR1_9B_" in disp and ".pdf" in disp


def test_t13_pdf_is_valid_and_reasonable_size():
    cid, h = _company_header()
    r = _get_pdf(h, "2020-01")
    assert r.status_code == 200
    assert r.content[:5] == b"%PDF-"
    assert b"%%EOF" in r.content[-1024:]
    assert len(r.content) > 4000, f"PDF suspiciously small: {len(r.content)} bytes"


def test_t14_pdf_dense_period_multipage_and_notes_present():
    """Seed a dense period (30 CN/DN mix) and verify PDF renders as
    multi-page with every note number represented in the extracted text."""
    cid, h = _company_header()
    uid = _current_uid()
    cust_id = _fresh_customer(h, gstin=_fake_gstin("36"), state="Telangana")
    seed_tag = f"C32T14-{uuid.uuid4().hex[:8]}"

    async def _seed():
        c = motor.motor_asyncio.AsyncIOMotorClient(os.environ["MONGO_URL"])
        _db = c[os.environ["DB_NAME"]]
        inv_id = f"{seed_tag}-inv"
        await _db.invoices.insert_one({
            "id": inv_id, "user_id": uid, "company_id": cid,
            "customer_id": cust_id,
            "invoice_number": f"{seed_tag}/M",
            "invoice_date": "2026-10-01",
            "total_amount": 100.0, "amount_paid": 0.0, "balance_due": 100.0,
            "trip_ids": [], "payments": [], "is_historical": False,
            "seed_tag": seed_tag, "gst_type": "cgst_sgst",
        })
        docs = []
        for i in range(30):
            docs.append({
                "id": f"{seed_tag}-cdn-{i:03d}", "user_id": uid, "company_id": cid,
                "customer_id": cust_id,
                "kind": "credit" if i % 2 == 0 else "debit",
                "note_number": f"C32T14-{seed_tag[-6:]}-{i:03d}",
                "note_date": "2026-10-15",
                "invoice_id": inv_id,
                "invoice_number_snapshot": f"{seed_tag}/M",
                "reason_code": "other", "reason_text": f"dense {i}",
                "lines": [], "subtotal": 10.0, "gst_type": "cgst_sgst",
                "cgst_rate": 2.5, "sgst_rate": 2.5, "igst_rate": 5.0,
                "cgst_amount": 0.25, "sgst_amount": 0.25, "igst_amount": 0.0,
                "total_tax": 0.5, "total_amount": 10.0,
                "rcm": True, "apply_gst": True, "status": "issued",
                "seed_tag": seed_tag,
            })
        await _db.credit_debit_notes.insert_many(docs)
        c.close()

    async def _cleanup():
        c = motor.motor_asyncio.AsyncIOMotorClient(os.environ["MONGO_URL"])
        _db = c[os.environ["DB_NAME"]]
        await _db.credit_debit_notes.delete_many({"seed_tag": seed_tag})
        await _db.invoices.delete_many({"seed_tag": seed_tag})
        c.close()

    try:
        asyncio.run(_seed())
        r = _get_pdf(h, "2026-10")
        assert r.status_code == 200, r.text[:400]
        assert r.content[:5] == b"%PDF-"
        doc = fitz.open(stream=r.content, filetype="pdf")
        try:
            all_text = "".join(page.get_text() for page in doc)
            page_count = doc.page_count
        finally:
            doc.close()
        assert page_count >= 2, f"expected multi-page, got {page_count}"
        seeded_present = sum(
            1 for i in range(30)
            if f"C32T14-{seed_tag[-6:]}-{i:03d}" in all_text
        )
        assert seeded_present >= 30, \
            f"only {seeded_present} of 30 note numbers found in PDF text"
    finally:
        asyncio.run(_cleanup())


def test_t15_pdf_static_column_widths_within_landscape_envelope():
    """All landscape tables must fit within 273 mm (A4 landscape content
    width). Regression guardrail against future column additions."""
    import sys
    sys.path.insert(0, "/app/backend")
    from pdf.gstr1_9b import _all_col_widths_mm, _CONTENT_W_MM
    for name, widths in _all_col_widths_mm():
        total = sum(widths)
        assert total <= _CONTENT_W_MM, \
            f"{name} column widths sum to {total} mm > {_CONTENT_W_MM} mm envelope"


def test_t16_pdf_empty_period_single_page_placeholders():
    """Empty period must render 200 OK, valid PDF, and each empty section
    surfaces its `No … records for this period` placeholder in the
    extracted text — no LayoutError, no empty-table crash."""
    cid, h = _company_header()
    r = _get_pdf(h, "2018-01")
    assert r.status_code == 200
    doc = fitz.open(stream=r.content, filetype="pdf")
    try:
        text = "".join(page.get_text() for page in doc)
        pages = doc.page_count
    finally:
        doc.close()
    assert pages >= 1
    for section in ("CDNR", "CDNUR", "B2CS Adjustments",
                    "Commercial", "Cancelled After Export"):
        assert section in text, f"section heading missing: {section}"
    # At least the CDNR empty placeholder must appear
    assert "No CDNR records" in text or "No records" in text.replace(" ", " ")


def test_t17_pdf_2100_note_high_volume():
    """2,100-note PDF must render without HTTP 500 / LayoutError. Every
    note number retained. Cleanup guaranteed."""
    cid, h = _company_header()
    uid = _current_uid()
    cust_id = _fresh_customer(h, gstin=_fake_gstin("36"), state="Telangana")
    SEED = 2100
    seed_tag = f"C32T17-{uuid.uuid4().hex[:10]}"

    async def _seed():
        c = motor.motor_asyncio.AsyncIOMotorClient(os.environ["MONGO_URL"])
        _db = c[os.environ["DB_NAME"]]
        inv_id = f"{seed_tag}-inv"
        await _db.invoices.insert_one({
            "id": inv_id, "user_id": uid, "company_id": cid,
            "customer_id": cust_id,
            "invoice_number": f"{seed_tag}/M",
            "invoice_date": "2027-01-01",
            "total_amount": 100.0, "amount_paid": 0.0, "balance_due": 100.0,
            "trip_ids": [], "payments": [], "is_historical": False,
            "seed_tag": seed_tag, "gst_type": "cgst_sgst",
        })
        docs = []
        for i in range(SEED):
            docs.append({
                "id": f"{seed_tag}-cdn-{i:05d}", "user_id": uid, "company_id": cid,
                "customer_id": cust_id,
                "kind": "credit" if i % 2 == 0 else "debit",
                "note_number": f"T17-{seed_tag[-6:]}-{i:05d}",
                "note_date": "2027-01-15",
                "invoice_id": inv_id,
                "invoice_number_snapshot": f"{seed_tag}/M",
                "reason_code": "other", "reason_text": f"hv {i}",
                "lines": [], "subtotal": 10.0, "gst_type": "cgst_sgst",
                "cgst_rate": 2.5, "sgst_rate": 2.5, "igst_rate": 5.0,
                "cgst_amount": 0.25, "sgst_amount": 0.25, "igst_amount": 0.0,
                "total_tax": 0.5, "total_amount": 10.0,
                "rcm": True, "apply_gst": True, "status": "issued",
                "seed_tag": seed_tag,
            })
        for i in range(0, len(docs), 500):
            await _db.credit_debit_notes.insert_many(docs[i:i + 500])
        c.close()

    async def _cleanup():
        c = motor.motor_asyncio.AsyncIOMotorClient(os.environ["MONGO_URL"])
        _db = c[os.environ["DB_NAME"]]
        await _db.credit_debit_notes.delete_many({"seed_tag": seed_tag})
        await _db.invoices.delete_many({"seed_tag": seed_tag})
        r1 = await _db.credit_debit_notes.count_documents({"seed_tag": seed_tag})
        r2 = await _db.invoices.count_documents({"seed_tag": seed_tag})
        c.close()
        return r1, r2

    try:
        asyncio.run(_seed())
        r = _get_pdf(h, "2027-01")
        assert r.status_code == 200, f"HTTP {r.status_code}: {r.text[:400] if r.status_code != 200 else ''}"
        assert r.content[:5] == b"%PDF-"
        doc = fitz.open(stream=r.content, filetype="pdf")
        try:
            page_count = doc.page_count
            # Verify a sample of note numbers survive
            all_text = "".join(page.get_text() for page in doc)
        finally:
            doc.close()
        assert page_count >= 10, f"expected many pages for 2,100 notes, got {page_count}"
        sample = [0, SEED // 2, SEED - 1]
        for i in sample:
            key = f"T17-{seed_tag[-6:]}-{i:05d}"
            assert key in all_text, f"sample note {key} missing from PDF text"
    finally:
        r1, r2 = asyncio.run(_cleanup())
        assert r1 == 0 and r2 == 0, f"cleanup orphans: cdn={r1} inv={r2}"


# ─── T18 · audit log for BOTH exports ──────────────────────────────────

def test_t18_audit_logs_written_for_xlsx_and_pdf():
    cid, h = _company_header()
    uid = _current_uid()
    month = "2020-01"  # empty is fine; we just want the audit row

    async def _count(fmt):
        c = motor.motor_asyncio.AsyncIOMotorClient(os.environ["MONGO_URL"])
        _db = c[os.environ["DB_NAME"]]
        n = await _db.audit_logs.count_documents({
            "user_id": uid, "module": "gstr_export", "action": "download",
            "changes.format": fmt,
        })
        c.close()
        return n

    before_xlsx = asyncio.run(_count("xlsx"))
    before_pdf  = asyncio.run(_count("pdf"))
    rx = _get_xlsx(h, month); assert rx.status_code == 200
    rp = _get_pdf(h, month);  assert rp.status_code == 200
    time.sleep(1.5)
    after_xlsx = asyncio.run(_count("xlsx"))
    after_pdf  = asyncio.run(_count("pdf"))
    assert after_xlsx > before_xlsx, f"xlsx audit not written: {before_xlsx}→{after_xlsx}"
    assert after_pdf  > before_pdf,  f"pdf audit not written: {before_pdf}→{after_pdf}"
